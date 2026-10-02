# -*- coding: utf-8 -*-
"""
把热搜追踪数据导出成甲方模板格式
================================================================================
数据源：data/hot_tracker.db（由 scripts/hot_tracker.py 持续写入）

输出（与甲方模板逐列一致）：
  事件数据集/
    事件.xlsx            24 列（列宽统一、行高自适应内容、自动换行）
    事件评论信息.xlsx     4 列（同上）
    事件图/{n}-{k}.jpg   事件配图，扁平存放：第 n 条事件的第 k 张

字段映射：
  序号            <- 行号
  事件名          <- 热搜词
  事件内容        <- 该话题下最热微博的正文
  来源            <- 微博账号昵称
  事件发生时间    <- 微博发布时间（= 事件最开始发生的时间）
  上头条热榜时间  <- 热搜接口的 onboard_time。**注意该字段的真实性质**：
                     它不是逐话题的真实秒级时刻，而是微博榜单的**批次刷新时刻**——
                     同一批上榜的话题共享同一值；且其"秒"部分呈**单调递增**
                     （实测 109 条里秒数从 13 一路递增到 24，真实时钟不可能这样）。
                     故精度约 ±1 分钟，同批次话题时间相同属数据源特性，非抓取错误。
  下头条热榜时间  <- 首次确认"已不在榜"的采样时刻，即下榜时刻的**上界**
                     （真实下榜落入 (最后一次在榜, 本值] 内，误差 <= 采样间隔）
  事件结束时间    <- 留空（微博数据不含此字段）。旧版直接复制下榜时间，
                     导致 69 行两列完全相同，已按反馈修正；
                     需要旧行为时加 --event-end offboard
  评论信息        <- R{序号}
  图片            <- P{序号}（需加 --fetch-pics 才会真正下载文件）
  类型            <- 微博官方分类
  是否为虚假      <- 留空（需人工核查）
  1小时~96小时    <- 上热搜后各时点的**热度值**（= 约定的浏览量口径）
                     掉榜取不到的时点留空，不伪造

关于图片（老师要求：每个事件 5 张，不足 5 张至少 1 张）：
  微博帖子本身就带图，接口能直接拿到原图 URL，所以**不必手工截图**，
  直接下载原图比截图更清晰。命名采用数字、扁平存放：
    事件图/{n}-{k}.jpg   第 n 条事件的第 k 张图（不再每个事件建一个子目录）
  表格「图片」列写 P{n}，即"该事件的图就是 {n}-* 那组"。
  · 默认不下载（导出快）；加 --fetch-pics 才下载
  · 帖子有 0 张图时（纯文字帖或视频帖），默认留空；
    加 --supplement-search 会去该话题搜索页第 1 页兜底收集图片，
    代价是**每个缺图事件多 1 次需要登录的搜索请求**，所以默认关闭

用法：
  python scripts/export_tracker.py                      # 只出表，不下载图片
  python scripts/export_tracker.py --fetch-pics         # 出表 + 下载配图
  python scripts/export_tracker.py --fetch-pics --supplement-search
  python scripts/export_tracker.py --no-zip
  python scripts/export_tracker.py --all-events         # 连还没抓正文的也导出
  python scripts/export_tracker.py --only-complete      # 只导出时间线完整的
"""
import argparse
import csv
import os
import re
import shutil
import sqlite3
import sys
import time
import urllib.parse
from datetime import datetime

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(ROOT, 'data', 'hot_tracker.db')
COOKIE_FILE = os.path.join(ROOT, 'config', 'cookie.txt')
OUTDIR = os.path.join(ROOT, 'dataset_tracker', '事件数据集')
IMGDIR = os.path.join(OUTDIR, '事件图')

UA_M = ('Mozilla/5.0 (iPhone; CPU iPhone OS 17_4 like Mac OS X) AppleWebKit/605.1.15 '
        '(KHTML, like Gecko) Version/17.4 Mobile/15E148 Safari/604.1')

OFFSETS = [1, 2, 4, 6, 8, 10, 12, 18, 24, 36, 48, 96]
OFFSET_COLS = ['1小时', '2小时', '4小时', '6小时', '8小时', '10小时',
               '12小时', '18小时', '24小时', '36小时', '48小时', '96小时']
EVENT_COLS = (['序号', '事件名', '事件内容', '来源', '事件发生时间',
               '上头条热榜时间', '下头条热榜时间', '事件结束时间',
               '评论信息', '图片', '类型', '是否为虚假'] + OFFSET_COLS)
COMMENT_COLS = ['ID', '评论者', '内容', '点赞数']


def parse_weibo_time(s):
    if not s:
        return ''
    try:
        return datetime.strptime(s, '%a %b %d %H:%M:%S %z %Y').strftime(
            '%Y-%m-%d %H:%M:%S')
    except ValueError:
        return s


def harvest_topic_pics(keyword, cookie, limit=5):
    """从该话题的搜索页第 1 页收集候选图片 URL。

    只在该事件自己没配图时兜底用。代价是 1 次需要登录的搜索请求，
    所以由 --supplement-search 显式开启，默认不跑。
    """
    from curl_cffi import requests as creq
    url = ('https://m.weibo.cn/api/container/getIndex?containerid='
           f'100103type%3D1%26q%3D{urllib.parse.quote(keyword)}'
           '&page_type=searchall&page=1')
    try:
        r = creq.get(url, headers={
            'User-Agent': UA_M, 'Referer': 'https://m.weibo.cn/',
            'Accept': 'application/json, text/plain, */*',
            'X-Requested-With': 'XMLHttpRequest', 'MWeibo-Pwa': '1',
            'Cookie': cookie}, impersonate='chrome', timeout=25)
        d = r.json()
    except Exception:                                # noqa: BLE001
        return []
    out = []
    for card in ((d.get('data') or {}).get('cards') or []):
        for cc in [card] + list(card.get('card_group') or []):
            m = cc.get('mblog')
            if not m:
                continue
            for p in (m.get('pics') or []):
                u = p.get('url') or (p.get('large') or {}).get('url')
                if u and u not in out:
                    out.append(u)
            if len(out) >= limit:
                return out[:limit]
    return out[:limit]


def _download_one(url, dest):
    """下一张图。优先 /mw690/（宽 690，清晰且体积合理），退 /large/，再退原始 URL。

    为什么不直接用 large：实测单张可达 3MB，一个事件 4 张就 11MB，
    一天几十个事件会把数据集撑爆；mw690 看图和放进报告都够用。
    """
    from curl_cffi import requests as creq
    cands = []
    for size in ('/mw690/', '/large/'):
        if '/orj360/' in url:
            cands.append(url.replace('/orj360/', size))
        elif '/thumb150/' in url:
            cands.append(url.replace('/thumb150/', size))
    cands.append(url)
    seen, ordered = set(), []
    for u in cands:
        if u not in seen:
            seen.add(u)
            ordered.append(u)
    for cand in ordered:
        try:
            r = creq.get(cand, headers={'User-Agent': UA_M,
                                        'Referer': 'https://weibo.com/'},
                         impersonate='chrome', timeout=30)
            if r.status_code == 200 and len(r.content) > 1200:
                with open(dest, 'wb') as fh:
                    fh.write(r.content)
                return len(r.content)
        except Exception:                            # noqa: BLE001
            continue
    return 0


def download_pics(targets, max_pics=5, supplement=False, verbose=True,
                  reuse_dir=None, download=True):
    """下载事件配图。

    踩过的坑（2026-09-22）：重导表格时若不带 --fetch-pics，"图片"列会**全空**——
    因为 pic_ref 只在下载流程里构建。而重下 125 张图要 15 分钟，纯属浪费。
    故加 reuse_dir：先把上一轮已下载的图挪到临时区，这里按 `{n}-*.jpg` 前缀
    直接匹配复用，秒级完成；只有缺失的才真正下载（download=False 则完全不下载）。

    targets = [(序号, event_id, word, pics_str, has_post),
               …]。has_post 用来区分两种情况：**还没抓正文**（后续轮次会补）
    和**帖子本身没配图**（本来就没有）—— 混为一谈会误导判断。

    命名：事件图/{n}-{k}.jpg —— 扁平存放，不再为每个事件建子目录。
    图片走的是新浪 CDN（http://wx*.sinaimg.cn），不是微博接口，
    所以间隔按 1.5s 起即可，不必像接口那样压到 10s 以上。
    """
    cookie = ''
    if os.path.exists(COOKIE_FILE):
        cookie = open(COOKIE_FILE, encoding='utf-8').read().strip()

    # 扫描上一轮已下载的图片，按事件序号归组
    cached = {}
    if reuse_dir and os.path.isdir(reuse_dir):
        for f in os.listdir(reuse_dir):
            m = re.match(r'^(\d+)-(\d+)\.jpg$', f)
            if m:
                cached.setdefault(int(m.group(1)), []).append(f)
        if verbose and cached:
            print(f'  复用上一轮配图：{len(cached)} 个事件有缓存')

    ref, n_ok, n_empty, n_nopost, n_fail, n_bytes = {}, 0, 0, 0, 0, 0
    n_reused = 0
    for idx, (i, eid, word, pics, has_post) in enumerate(targets, 1):
        # ① 优先复用缓存
        if i in cached:
            for f in sorted(cached[i]):
                try:
                    shutil.copy2(os.path.join(reuse_dir, f),
                                 os.path.join(IMGDIR, f))
                    n_bytes += os.path.getsize(os.path.join(IMGDIR, f))
                except Exception:                               # noqa: BLE001
                    pass
            ref[eid] = f'P{i}'
            n_ok += 1
            n_reused += 1
            continue
        # ② 没有缓存：只在允许下载时才走网络
        if not download:
            if not has_post:
                n_nopost += 1
            else:
                n_empty += 1
            continue

        urls = [u for u in (pics or '').split('|') if u][:max_pics]
        src = '帖子自带'
        if not urls and supplement and cookie:
            urls = harvest_topic_pics(word, cookie, limit=max_pics)
            src = '话题搜索兜底'
            time.sleep(2.0)
        if not urls:
            if not has_post:
                # 注意别把这两种情况混为一谈：
                # 还没抓正文 ≠ 帖子没配图。前者靠后续轮次补，后者本来就没有。
                n_nopost += 1
                if verbose:
                    print(f'    P{i} 尚未抓正文（详情限流，待后续轮次）  '
                          f'{(word or "")[:20]}')
            else:
                n_empty += 1
                if verbose:
                    print(f'    P{i} 帖子本身未配图  {(word or "")[:20]}')
            continue
        sub = IMGDIR
        got = 0
        for k, u in enumerate(urls, 1):
            n = _download_one(u, os.path.join(sub, f'{i}-{k}.jpg'))
            if n:
                got += 1
                n_bytes += n
            time.sleep(1.5)
        if got:
            ref[eid] = f'P{i}'
            n_ok += 1
            if verbose:
                flag = '' if got >= max_pics else f'（不足 {max_pics} 张，帖子里只有这些）'
                print(f'    P{i} {got} 张  [{src}]  {(word or "")[:20]} {flag}')
        else:
            n_fail += 1
            if verbose:
                print(f'    P{i} 下载失败  {(word or "")[:20]}')
    print(f'  图片：成功 {n_ok} 个事件（其中复用缓存 {n_reused} 个），'
          f'共 {n_bytes / 1048576:.1f} MB；'
          f'帖子本身无图 {n_empty} 个；尚未抓正文 {n_nopost} 个；失败 {n_fail} 个')
    return ref


def _disp_len(s):
    """显示宽度：中日韩/全角字符算 2，其余算 1（Excel 列宽以半角字符计）。"""
    n = 0
    for ch in str(s):
        n += 2 if ord(ch) > 0x2E80 else 1
    return n


def _wrapped_lines(text, col_width):
    """文本在给定列宽下换行后占的行数。"""
    if text is None or text == '':
        return 1
    total = 0
    for para in str(text).split('\n'):
        total += max(1, -(-_disp_len(para) // max(1, int(col_width))))
    return total


def style_sheet(ws, widths, line_pt=14.5, min_h=18.0, max_h=409.0):
    """统一列宽 + 行高按内容自适应 + 自动换行，保证格子能显示完整内容。

    · 每列固定一个宽度（同一列所有格子左右宽度一致）
    · 每行高度按该行内容自动撑开（不裁切文字）
    · 全部单元格开启自动换行、顶端对齐；表头冻结并加底色
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ncol, nrow = len(widths), ws.max_row
    # 1) 列宽（每列统一）
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # 2) 表头
    head_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, size=11)
        cell.fill = PatternFill('solid', fgColor='D9E1F2')
        cell.alignment = head_align
    ws.row_dimensions[1].height = 26
    # 3) 数据行：行高按内容自适应
    body_align = Alignment(vertical='top', wrap_text=True)
    for r in range(2, nrow + 1):
        need = 1
        for c in range(1, ncol + 1):
            need = max(need, _wrapped_lines(ws.cell(r, c).value, widths[c - 1]))
        ws.row_dimensions[r].height = min(max_h, max(min_h, need * line_pt))
        for c in range(1, ncol + 1):
            ws.cell(r, c).alignment = body_align
    # 4) 冻结表头，滚动时列名常驻
    ws.freeze_panes = 'A2'


# 各列宽度（半角字符为单位；中文按 2 计）
# 「事件内容」列最宽：实测最长正文显示宽度 2424，列宽需 ≥87 才能让所有内容
# 在 Excel 单行 409pt 上限内完整显示（宽 90 时最长行 392pt，安全）。
EVENT_WIDTHS = ([6, 24, 90, 16, 20, 20, 20, 20, 10, 10, 10, 12]
                + [10] * len(OFFSET_COLS))
COMMENT_WIDTHS = [8, 20, 90, 10]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--no-zip', action='store_true')
    ap.add_argument('--only-complete', action='store_true',
                    help='只导出有正文且至少有 2 个热度时点的')
    ap.add_argument('--all-events', action='store_true',
                    help='连"还没抓到正文"的话题也一起导出（默认只导出有正文的，'
                         '避免表里出现成片空白）')
    ap.add_argument('--fetch-pics', action='store_true',
                    help='下载事件配图到 事件图/{n}-{k}.jpg（不指定则只出表）')
    ap.add_argument('--reuse-pics', action='store_true',
                    help='复用上一轮已下载的配图来填「图片」列（不重新下载）。'
                         '重导表格时务必带上，否则「图片」列会全空')
    ap.add_argument('--max-pics', type=int, default=5,
                    help='每个事件最多存几张图（默认 5，与老师要求一致）')
    ap.add_argument('--supplement-search', action='store_true',
                    help='帖子自己没配图时，去该话题搜索页兜底找图。'
                         '每个缺图事件多 1 次需要登录的搜索请求，默认关闭')
    ap.add_argument('--event-end', choices=['blank', 'offboard'], default='blank',
                    help='「事件结束时间」列怎么填：blank=留空（默认，微博无此数据）；'
                         'offboard=与下头条热榜时间相同（旧行为，会造成两列完全重复）')
    args = ap.parse_args()

    if not os.path.exists(DB):
        print('未找到 data/hot_tracker.db，请先跑 scripts/hot_tracker.py --tick')
        return 1
    import openpyxl

    # 要用上一轮已下载的图时，必须先把 事件图 挪出 dataset_tracker，
    # 否则会被下面的 rmtree 一并删掉（重下 125 张图需要约 15 分钟）。
    #
    # ⚠️ 2026-09-22 踩坑：最初写成「old_pic 存在就 rmtree(stash) 再 move」，
    # 但崩溃重跑时会留下一个**空的事件图目录**，于是 rmtree(stash) 把上一轮
    # 暂存的好图全删了。现在改为：**stash 里已有图就绝不删**，只有「stash 为空
    # 且 old_pic 确实有图」时才搬运。宁可不复用，也不能删掉已有素材。
    stash = os.path.join(ROOT, '_pic_stash')
    if args.reuse_pics:
        old_pic = os.path.join(ROOT, 'dataset_tracker', '事件数据集', '事件图')
        old_n = len(os.listdir(old_pic)) if os.path.isdir(old_pic) else 0
        stash_n = len(os.listdir(stash)) if os.path.isdir(stash) else 0
        if stash_n:
            print(f'  复用已有暂存配图：{stash_n} 个文件（不重新搬运/删除）')
        elif old_n:
            try:
                shutil.move(old_pic, stash)
                print(f'  已暂存上一轮配图：{old_n} 个文件')
            except Exception as exc:                            # noqa: BLE001
                print(f'  暂存配图失败：{exc}')

    shutil.rmtree(os.path.join(ROOT, 'dataset_tracker'), ignore_errors=True)
    os.makedirs(IMGDIR, exist_ok=True)

    c = sqlite3.connect(DB)
    rows = c.execute("""
        SELECT e.event_id, e.word, e.category, e.onboard_ts, e.offboard_ts,
               p.created_at, p.text, p.screen_name, p.mid, p.pics,
               e.offboard_upper_ts
        FROM tracked_events e
        LEFT JOIN event_posts p ON p.event_id = e.event_id
        ORDER BY e.onboard_ts DESC
    """).fetchall()

    if args.only_complete:
        keep = []
        for r in rows:
            n = c.execute("SELECT COUNT(*) FROM heat_series "
                          "WHERE event_id=? AND status='recorded'",
                          (r[0],)).fetchone()[0]
            if r[6] and n >= 2:
                keep.append(r)
        rows = keep

    # 默认只导出**已经有正文**的事件。
    # 为什么：热搜榜上大量话题只是短暂上榜、还没轮到抓正文（详情有限流），
    # 如果把它们一并导出，表里会出现成片空白的「事件内容/来源/图片/评论信息」，
    # 交付出去像是没做完。老师的口径也是"一天约 40 个事件"，
    # 所以只交"抓全了"的那些，看的是质量不是行数。
    # 需要看全部追踪对象时加 --all-events。
    if not args.all_events:
        before = len(rows)
        rows = [r for r in rows if r[6]]
        if before != len(rows):
            print(f'  只导出已有正文的事件：{len(rows)} / {before} 条'
                  f'（其余 {before - len(rows)} 条还没轮到抓正文，'
                  f'用 --all-events 可一并导出）')

    # 先决定序号 -> 下载图片（序号必须与表里的"序号"列一致，所以要在构表前做）
    if args.fetch_pics or args.reuse_pics:
        targets = [(i, r[0], r[1], r[9], bool(r[8] or r[6]))
                   for i, r in enumerate(rows, 1)]
        pic_ref = download_pics(targets, max_pics=args.max_pics,
                                supplement=args.supplement_search,
                                reuse_dir=(stash if args.reuse_pics else None),
                                download=args.fetch_pics)
    else:
        pic_ref = {}

    event_rows, comment_rows = [], []
    n_with_heat = 0
    for i, (eid, word, cat, onb, off, created, text, sn, mid,
            _pics, off_up) in enumerate(rows, 1):
        series = dict(c.execute(
            "SELECT offset_hours, heat FROM heat_series "
            "WHERE event_id=? AND status='recorded'", (eid,)).fetchall())
        if series:
            n_with_heat += 1
        heat_cells = [series.get(o) if series.get(o) is not None else ''
                      for o in OFFSETS]
        # ── 下头条热榜时间 ──────────────────────────────────────────────
        # 取「首次确认不在榜」的上界（offboard_upper_ts），而不是「最后一次
        # 在榜」的下界。原因：抽样间隔 >= 30 分钟，只被采到一轮的短命话题
        # 若取下界，会显示成「上榜 18 秒就下榜」——那是观测假象，不是事实。
        # 真实下榜时刻落在 (下界, 上界] 内；取上界可保证不出现自相矛盾的值。
        off_best = off_up or off
        # ── 事件结束时间 ────────────────────────────────────────────────
        # 默认**留空**：微博数据里没有「事件结束」这个字段，过去直接复制
        # 下榜时间，导致 69 行两列完全相同，属误导。需要旧行为时加
        # --event-end offboard。
        end_cell = off_best if args.event_end == 'offboard' else ''
        event_rows.append([
            i, word or '', text or '', sn or '', parse_weibo_time(created),
            onb or '', off_best or '', end_cell,
            f'R{i}', pic_ref.get(eid, ''), cat or '', '',
        ] + heat_cells)

        # 评论：内容 + 点赞数（不含时间），按点赞降序
        for cm in c.execute("""SELECT screen_name, text, like_count
                               FROM event_comments WHERE event_id=?
                               ORDER BY COALESCE(like_count,0) DESC""", (eid,)):
            comment_rows.append([f'R{i}', cm[0] or '', cm[1] or '',
                                 cm[2] if cm[2] is not None else 0])
    c.close()

    print(f'事件 {len(event_rows)} 条（其中有热度轨迹的 {n_with_heat} 条），'
          f'评论 {len(comment_rows)} 条')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(EVENT_COLS)
    for r in event_rows:
        ws.append(r)
    style_sheet(ws, EVENT_WIDTHS)          # 列宽统一 + 行高自适应 + 自动换行
    p1 = os.path.join(OUTDIR, '事件.xlsx')
    wb.save(p1)

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = 'Sheet1'
    ws2.append(COMMENT_COLS)
    for r in comment_rows:
        ws2.append(r)
    style_sheet(ws2, COMMENT_WIDTHS)       # 同上
    p2 = os.path.join(OUTDIR, '事件评论信息.xlsx')
    wb2.save(p2)

    for name, cols, data in (('事件.csv', EVENT_COLS, event_rows),
                             ('事件评论信息.csv', COMMENT_COLS, comment_rows)):
        with open(os.path.join(OUTDIR, name), 'w', encoding='utf-8-sig',
                  newline='') as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            w.writerows(data)

    print(f'✓ {p1}')
    print(f'✓ {p2}')
    n_img = sum(len(fs) for _, _, fs in os.walk(IMGDIR))
    print(f'  事件图/ 共 {n_img} 个文件'
          + ('' if n_img else '（未下载图片，加 --fetch-pics）'))
    if not args.no_zip:
        zp = os.path.join(ROOT, 'dataset_tracker', '事件数据集.zip')
        shutil.make_archive(zp[:-4], 'zip',
                            os.path.join(ROOT, 'dataset_tracker'), '事件数据集')
        print(f'✓ {zp}  ({os.path.getsize(zp):,} bytes)')
    # 清理暂存的上一轮配图
    if os.path.isdir(stash):
        shutil.rmtree(stash, ignore_errors=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
